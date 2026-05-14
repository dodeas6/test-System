import streamlit as st
import pandas as pd
import io
import re
from datetime import datetime
from openpyxl import load_workbook

LOGO_URL = "https://i.pinimg.com/originals/5a/65/ee/5a65ee278cd557143f05a4ba91abbfa8.gif"
DEEP_GREEN_BG = "#021a0d"  
BUTTON_GREEN = "#16a34a" 

PASSENGER_ICON_URL = "https://i.pinimg.com/1200x/e2/ed/71/e2ed710d953174362c9510d327d04d53.jpg" 
CREW_ICON_URL = "https://i.pinimg.com/736x/3d/2c/c4/3d2cc44e42cb3898bae3cbdc3bcc3e1a.jpg"

st.set_page_config(page_title="Iraqi Airways System", layout="wide")

# هندسة التصميم بالـ CSS مع تصفير إطار وألوان خلفية أزرار اللغة لتصبح داكنة كلياً
st.markdown(f"""
    <style>
    @import url('https://googleapis.com');
    .stApp {{ background-color: {DEEP_GREEN_BG} !important; }}
    .stMarkdown, p, label, h3, h2, h1, span {{ color: white !important; font-family: 'Cairo', sans-serif; }}
    .header-content div {{ color: #1d4c2b !important; font-weight: bold !important; display: block !important; }}
    [data-testid="stFileUploader"] {{ background-color: {BUTTON_GREEN} !important; border-radius: 12px !important; padding: 15px !important; }}
    [data-testid="stFileUploader"] * {{ color: black !important; font-weight: 900 !important; }}
    .stTextArea textarea {{ background-color: #dcfce7 !important; border: 2px solid #4ade80 !important; border-radius: 15px !important; color: #013220 !important; font-weight: bold; }}
    .stTextInput input {{ background-color: #ffffff !important; color: #000000 !important; border: 2px solid #16a34a !important; border-radius: 8px !important; font-weight: bold !important; }}
    div.stFormSubmitButton {{ text-align: center !important; display: flex !important; justify-content: center !important; margin-top: 25px !important; }}
    div.stFormSubmitButton > button {{ width: 280px !important; background-color: {BUTTON_GREEN} !important; color: white !important; border-radius: 8px !important; font-weight: 900 !important; padding: 12px 24px !important; border: none !important; }}
    .stButton>button, .stDownloadButton>button {{ width: 100% !important; background-color: {BUTTON_GREEN} !important; color: white !important; border-radius: 12px !important; font-weight: 900 !important; padding: 15px !important; border: none !important; }}
    
    /* فرض أزرار اللغة لتكون داكنة ومدمجة بالكامل مع لون خلفية النظام والشاشات */
    div.lang-block div[data-testid="stColumn"] button {{
        background-color: {DEEP_GREEN_BG} !important;
        color: white !important;
        border: 1px solid rgba(255,255,255,0.2) !important;
        font-size: 14px !important;
        font-weight: bold !important;
        box-shadow: none !important;
    }}
    div.lang-block div[data-testid="stColumn"] button:hover {{
        border-color: {BUTTON_GREEN} !important;
        color: {BUTTON_GREEN} !important;
    }}
    </style>
    """, unsafe_allow_html=True)

trans = {
    "AR": {
        "m_title": "منظومة الخطوط الجوية العراقية الموحدة",
        "p_title": "نظام إدارة بيانات المسافرين والحقائب",
        "c_title": "نظام إدارة بيانات الطواقم",
        "sel_sys": "الرجاء اختيار النظام المطلوب لبدء العمل:",
        "gate": "بوابة دخول الموظفين",
        "user": "اسم المستخدم",
        "pass": "كلمة المرور",
        "login": "دخول",
        "err_login": "خطأ في البيانات",
        "open_p": "فتح نظام المسافرين والحقائب (المانيفست)",
        "open_c": "فتح نظام إدارة بيانات الطواقم (إدخال يدوي)",
        "back": "⬅️ العودة للقائمة الرئيسية",
        "up_lbl": "📄 الخطوة الأولى: رفع ملف قالب الإكسل (template.xlsx)",
        "p_paste": "📋 الصق بيانات المانيفست هنا:",
        "b_paste": "👜 الصق بيانات الحقائب هنا:",
        "proc": "🚀 معالجة البيانات وتوليد التقرير",
        "err_tmp": "يرجى رفع ملف template.xlsx اولاً",
        "ok_match": "✅ تمت مطابقة {} من الأفراد.",
        "dl_lbl": "📥 تحميل التقرير النهائي المكتمل",
        "entry_hdr": "📝 إدخال بيانات فرد جديد من الطاقم",
        "edit_hdr": "📝 تعديل بيانات فرد من الطاقم (رقم: {})",
        "add_b": "➕ إضافة فرد طاقم",
        "save_b": "💾 حفظ التعديلات",
        "cancel_b": "❌ إلغاء التعديل",
        "err_ar": "⚠️ خطأ: لا يمكن إدخال الحروف العربية. يرجى ملء الحقول بالإنجليزية فقط!",
        "err_mand": "⚠️ يرجى ملء الحقول الأساسية على الأقل (الاسم، اللقب، ورقم الجواز).",
        "ok_save": "✅ تم حفظ البيانات بنجاح!",
        "curr_hdr": "📋 قائمة الأفراد المضافة حالياً",
        "act_hdr": "🛠️ الإجراءات المتاحة:",
        "edit_act": "✏️ تعديل",
        "del_act": "❌ حذف",
        "empty": "💡 الجدول فارغ حالياً. قم بإضافة أفراد عبر الاستمارة في الأعلى.",
        "f1": "1. الاسم (Name)", "f2": "2. اللقب (Surname)",
        "f3": "3. المواليد (Date of Birth)", "f4": "4. الجنسية (Nationality)",
        "f5": "5. درجة الطاقم (Type Of Crew)", "f6": "6. نوع الجواز (Document Type)",
        "f7": "7. رقم الجواز (Document Number)"
    },
    "EN": {
        "m_title": "Iraqi Airways Unified System",
        "p_title": "Passengers & Baggage Management System",
        "c_title": "Crew Management System",
        "sel_sys": "Please select the required system to start working:",
        "gate": "Staff Login Gateway",
        "user": "Username", "pass": "Password", "login": "Login",
        "err_login": "Invalid credentials",
        "open_p": "Open Passengers & Baggage System (Manifest)",
        "open_c": "Open Crew Management System (Manual Entry)",
        "back": "⬅️ Back to Main Menu",
        "up_lbl": "📄 Step 1: Upload Excel Template (template.xlsx)",
        "p_paste": "📋 Paste Manifest Data Here:",
        "b_paste": "👜 Paste Baggage Data Here:",
        "proc": "🚀 Process Data & Generate Report",
        "err_tmp": "Please upload the template.xlsx file first.",
        "ok_match": "✅ Successfully processed {} members.",
        "dl_lbl": "📥 Download Completed Final Report",
        "entry_hdr": "📝 Crew Member Data Entry",
        "edit_hdr": "📝 Edit Crew Member Data (Index: {})",
        "add_b": "➕ Add Crew Member", "save_b": "💾 Save Changes", "cancel_b": "❌ Cancel Edit",
        "err_ar": "⚠️ Error: Arabic characters are not allowed. Please fill all fields in English only!",
        "err_mand": "⚠️ Please fill in at least the mandatory fields: Name, Surname, and Document Number.",
        "ok_save": "✅ Data saved successfully!",
        "curr_hdr": "📋 Current Members List",
        "act_hdr": "🛠️ Actions Available:",
        "edit_act": "✏️ Edit", "del_act": "❌ Delete",
        "empty": "💡 No members added yet. Fill out the form above to populate the table.",
        "f1": "1. Name", "f2": "2. Surname", "f3": "3. Date of Birth", "f4": "4. Nationality",
        "f5": "5. Type Of Crew (REF)", "f6": "6. Document Type(REF)", "f7": "7. Document Number"
    }
}

def format_dob(dob_str):
    if not dob_str: return ""
    clean_str = re.sub(r'\s+', '', dob_str)
    parts = re.split(r'[-/.]', clean_str)
    if len(parts) == 3:
        try:
            year, month, day = int(parts[0]), int(parts[1]), int(parts[2])
            if year < 100: day, month, year = int(parts[0]), int(parts[1]), int(parts[2])
            dt = datetime(year, month, day)
            return dt.strftime("%d%b%Y")
        except ValueError: return dob_str.upper()
    return dob_str.upper()

def contains_arabic(text):
    if not text: return False
    return bool(re.search(r'[\u0600-\u06FF]', text))

def extract_val(text, start_key):
    match = re.search(rf"{start_key}-(.*?)(?=/|$)", text, re.IGNORECASE)
    return match.group(1).strip() if match else ""

def find_seat(text):
    match = re.search(r'\b0*(\d{1,3}[A-F])\b', text, re.IGNORECASE)
    return match.group(1).upper() if match else None

def parse_bags(text):
    pcs = re.search(r'(\d+)\s*PCS', text, re.IGNORECASE)
    kg = re.search(r'(\d+)\s*KG', text, re.IGNORECASE)
    return (pcs.group(1) if pcs else "0"), (kg.group(1) if kg else "0")

def render_header(system_title):
    st.markdown(f"""
    <div style="background-color:white; padding:30px; border-radius:0 0 35px 35px; text-align:center; box-shadow: 0 10px 20px rgba(0,0,0,0.5);" class="header-content">
        <center><img src="{LOGO_URL}" width="180"></center>
        <div style="font-size:32px; margin-top:15px; color:#1d4c2b !important; font-weight:900;">{system_title}</div>
        <div style="font-family:Orbitron; font-size:24px; letter-spacing:4px; margin-top:5px; color:#1d4c2b !important; font-weight:bold;">IRAQI AIRWAYS</div>
        <div style="font-family:Cairo; font-size:18px; margin-top:5px; color:#16a34a !important; font-weight:bold;">Designed by Eng. Mohammad Abdulsattar</div>
    </div>
    <br>
    """, unsafe_allow_html=True)

if "auth" not in st.session_state: st.session_state["auth"] = False
if "current_page" not in st.session_state: st.session_state["current_page"] = "dashboard" 
if "crew_list" not in st.session_state: st.session_state["crew_list"] = []
if "edit_index" not in st.session_state: st.session_state["edit_index"] = None
if "focus_flag" not in st.session_state: st.session_state["focus_flag"] = False
if "lang" not in st.session_state: st.session_state["lang"] = "AR"

curr_lang = st.session_state["lang"]

if not st.session_state["auth"]:
    st.markdown(f'<div style="text-align:center; padding-top:40px;"><img src="{LOGO_URL}" width="230"></div>', unsafe_allow_html=True)
    _, col_mid, _ = st.columns([1, 2, 1])
    with col_mid:
        st.markdown(f"<h2 style='text-align:center;'>{trans[curr_lang]['gate']}</h2>", unsafe_allow_html=True)
        user = st.text_input(trans[curr_lang]["user"])
        pw = st.text_input(trans[curr_lang]["pass"], type="password")
        if st.button(trans[curr_lang]["login"]):
            if user == "iraqi_admin" and pw == "ia2024":
                st.session_state["auth"] = True
                st.rerun()
            else: st.error(trans[curr_lang]["err_login"])
        
        st.markdown("<div style='text-align:center; font-size:16px; color:#16a34a; font-weight:bold; margin-top:15px;'>Designed by Eng. Mohammad Abdulsattar</div><br>", unsafe_allow_html=True)
        
        st.markdown("<div class='lang-block'>", unsafe_allow_html=True)
        l_c1, l_c2 = st.columns(2)
        with l_c1:
            if st.button("العربية", key="btn_ar"):
                st.session_state["lang"] = "AR"
                st.rerun()
        with l_c2:
            if st.button("English", key="btn_en"):
                st.session_state["lang"] = "EN"
                st.rerun()
        st.markdown("</div>", unsafe_allow_html=True)
    st.stop()

# لوحة التحكم المركزية
if st.session_state["current_page"] == "dashboard":
    render_header(trans[curr_lang]["m_title"])
    st.markdown(f"<h3 style='text-align:center; font-weight: bold;'>{trans[curr_lang]['sel_sys']}</h3><br>", unsafe_allow_html=True)
    _, col_left, _, col_right, _ = st.columns([1, 2, 1, 2, 1])
    with col_left:
        st.markdown(f'<div style="text-align:center; background-color:white; border-radius:20px; padding:20px; border:3px solid #16a34a;"><img src="{PASSENGER_ICON_URL}" style="width:130px; height:130px; border-radius:12px; object-fit:cover; margin-bottom:10px;"><div style="color:#1d4c2b; font-size:20px; font-weight:900;">Passengers System</div></div><br>', unsafe_allow_html=True)
        if st.button(trans[curr_lang]["open_p"], key="dash_p_btn"):
            st.session_state["current_page"] = "passengers"
            st.rerun()
    with col_right:
        st.markdown(f'<div style="text-align:center; background-color:white; border-radius:20px; padding:20px; border:3px solid #16a34a;"><img src="{CREW_ICON_URL}" style="width:130px; height:130px; border-radius:12px; object-fit:cover; margin-bottom:10px;"><div style="color:#1d4c2b; font-size:20px; font-weight:900;">Crew System</div></div><br>', unsafe_allow_html=True)
        if st.button(trans[curr_lang]["open_c"], key="dash_c_btn"):
            st.session_state["current_page"] = "crew"
            st.rerun()

# نظام الركاب والحقائب الأصلي المترجم
elif st.session_state["current_page"] == "passengers":
    render_header(trans[curr_lang]["p_title"])
    if st.button(trans[curr_lang]["back"], key="back_p"):
        st.session_state["current_page"] = "dashboard"
        st.rerun()
    template_file = st.file_uploader(trans[curr_lang]["up_lbl"], type=["xlsx"], key="file_p")
    c1, c2 = st.columns(2)
    with c1: p_input = st.text_area(trans[curr_lang]["p_paste"], height=280, key="txt_p1")
    with c2: b_input = st.text_area(trans[curr_lang]["b_paste"], height=280, key="txt_p2")
    if st.button(trans[curr_lang]["proc"], key="process_p"):
        if not template_file: st.error(trans[curr_lang]["err_tmp"])
        elif p_input and b_input:
            p_sections = re.split(r'(\d+\.)', p_input)
            passengers = {}
            for i in range(1, len(p_sections), 2):
                entry = p_sections[i] + p_sections[i+1]
                seat = find_seat(entry)
                if seat:
                    passengers[seat] = {
                        "first": extract_val(entry, "FIRST NAME"), "last": extract_val(entry, "SURNAME"),
                        "nat": extract_val(entry, "NATIONALITY"), "gender": extract_val(entry, "GENDER"),
                        "passport": extract_val(entry, "NUMBER"), "seat": seat, "pcs": "0", "kg": "0"
                    }
            for row in b_input.split('\n'):
                b_seat = find_seat(row)
                if b_seat and b_seat in passengers:
                    pcs, kg = parse_bags(row)
                    passengers[b_seat]["pcs"], passengers[b_seat]["kg"] = pcs, kg
            if passengers:
                st.success(trans[curr_lang]["ok_match"].format(len(passengers)))
                st.dataframe(pd.DataFrame(list(passengers.values())), use_container_width=True)
                wb = load_workbook(template_file)
                ws = wb.active
                for i, data in enumerate(list(passengers.values())):
                    curr_row = 2 + i 
                    ws.cell(row=curr_row, column=1).value = data["first"]
                    ws.cell(row=curr_row, column=2).value = data["last"]
                    ws.cell(row=curr_row, column=3).value = data["nat"]
                    ws.cell(row=curr_row, column=4).value = data["gender"]
                    ws.cell(row=curr_row, column=5).value = "PASSPORT"
                    ws.cell(row=curr_row, column=6).value = data["passport"]
                    ws.cell(row=curr_row, column=7).value = data["seat"]
                    ws.cell(row=curr_row, column=8).value = data["pcs"]
                    ws.cell(row=curr_row, column=9).value = data["kg"]
                out = io.BytesIO()
                wb.save(out)
                st.download_button(trans[curr_lang]["dl_lbl"], out.getvalue(), "Iraqi_Airways_Final_Report.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

# نظام الطواقم مع التعريب الشامل للحقول
elif st.session_state["current_page"] == "crew":
    render_header(trans[curr_lang]["c_title"])
    if st.button(trans[curr_lang]["back"], key="back_c"):
        st.session_state["current_page"] = "dashboard"
        st.rerun()
    template_file = st.file_uploader(trans[curr_lang]["up_lbl"], type=["xlsx"], key="file_c")
    edit_idx = st.session_state["edit_index"]
    
    if edit_idx is not None:
        st.markdown(f"### {trans[curr_lang]['edit_hdr'].format(edit_idx + 1)}")
        current_data = st.session_state["crew_list"][edit_idx]
        default_name = current_data["Name"]
        default_surname = current_data["Surname"]
        default_dob = current_data["Date of Birth"]
        default_nat = current_data["Nationality"]
        default_crew_type = current_data["Type Of Crew (REF)"]
        default_doc_type = current_data["Document Type(REF)"]
        default_doc_num = current_data["Document Number"]
        btn_label = trans[curr_lang]["save_b"]
    else:
        st.markdown(f"### {trans[curr_lang]['entry_hdr']}")
        default_name, default_surname, default_dob, default_nat = "", "", "", ""
        default_crew_type, default_doc_type, default_doc_num = "", "", ""
        btn_label = trans[curr_lang]["add_b"]

    if st.session_state["focus_flag"]:
        st.markdown("""<script>var inputs = window.parent.document.getElementsByTagName('input');for (var i = 0; i < inputs.length; i++) {if (inputs[i].getAttribute('aria-label') === '1. Name' || inputs[i].getAttribute('aria-label') === '1. الاسم (Name)') {inputs[i].focus();break;}}</script>""", unsafe_allow_html=True)
        st.session_state["focus_flag"] = False

    with st.form("crew_form", clear_on_submit=False):
        row1_col1, row1_col2 = st.columns(2)
        with row1_col1: v_name = st.text_input(trans[curr_lang]["f1"], value=default_name, key=f"n_{edit_idx}")
        with row1_col2: v_surname = st.text_input(trans[curr_lang]["f2"], value=default_surname, key=f"s_{edit_idx}")
        row2_col1, row2_col2 = st.columns(2)
        with row2_col1: v_dob = st.text_input(trans[curr_lang]["f3"], value=default_dob, placeholder="e.g., 1992-6-22", key=f"d_{edit_idx}")
        with row2_col2: v_nat = st.text_input(trans[curr_lang]["f4"], value=default_nat, key=f"na_{edit_idx}")
        row3_col1, row3_col2 = st.columns(2)
        with row3_col1: v_crew_type = st.text_input(trans[curr_lang]["f5"], value=default_crew_type, key=f"c_{edit_idx}")
        with row3_col2: v_doc_type = st.text_input(trans[curr_lang]["f6"], value=default_doc_type, key=f"dt_{edit_idx}")
        row4_col1, row4_col2 = st.columns(2)
        with row4_col1: v_doc_num = st.text_input(trans[curr_lang]["f7"], value=default_doc_num, key=f"dn_{edit_idx}")
        with row4_col2: st.write("") 
        submit_btn = st.form_submit_button(btn_label)

    if edit_idx is not None and st.button(trans[curr_lang]["cancel_b"], key="cancel_c"):
        st.session_state["edit_index"] = None
        st.session_state["focus_flag"] = False
        st.rerun()

    if submit_btn:
        all_fields = [v_name, v_surname, v_dob, v_nat, v_crew_type, v_doc_type, v_doc_num]
        if any(contains_arabic(f) for f in all_fields): st.error(trans[curr_lang]["err_ar"])
        elif v_name and v_surname and v_doc_num:
            formatted_dob_val = format_dob(v_dob)
            new_crew = {
                "Name": v_name.strip().upper(), "Surname": v_surname.strip().upper(),
                "Date of Birth": formatted_dob_val, "Nationality": v_nat.strip().upper(),
                "Type Of Crew (REF)": v_crew_type.strip().upper(), "Document Type(REF)": v_doc_type.strip().upper(),
                "Document Number": v_doc_num.strip().upper()
            }
            if edit_idx is not None:
                st.session_state["crew_list"][edit_idx] = new_crew
                st.session_state["edit_index"] = None
                st.session_state["focus_flag"] = False
            else: st.session_state["crew_list"].append(new_crew)
            st.success(trans[curr_lang]["ok_save"])
            st.rerun()
        else: st.error(trans[curr_lang]["err_mand"])

    if st.session_state["crew_list"]:
        st.markdown(f"### {trans[curr_lang]['curr_hdr']}")
        st.dataframe(pd.DataFrame(st.session_state["crew_list"]), use_container_width=True)
        st.markdown(f"**{trans[curr_lang]['act_hdr']}**")
        for idx, member in enumerate(st.session_state["crew_list"]):
            act_col1, act_col2, act_col3 = st.columns([3, 1, 1])
            with act_col1: st.write(f"👤 {member['Name']} {member['Surname']} [{member['Document Number']}]")
            with act_col2:
                if st.button(trans[curr_lang]["edit_act"], key=f"edit_{idx}", use_container_width=True):
                    st.session_state["edit_index"] = idx
                    st.session_state["focus_flag"] = True
                    st.rerun()
            with act_col3:
                if st.button(trans[curr_lang]["del_act"], key=f"del_{idx}", use_container_width=True):
                    st.session_state["crew_list"].pop(idx)
                    if st.session_state["edit_index"] == idx:
                        st.session_state["edit_index"] = None
                        st.session_state["focus_flag"] = False
                    st.rerun()
        st.markdown("---")
        if st.button(trans[curr_lang]["proc"], use_container_width=True, key="generate_c"):
            if not template_file: st.error(trans[curr_lang]["err_tmp"])
            else:
                try:
                    wb = load_workbook(template_file)
                    ws = wb.active
                    for row_idx, data in enumerate(st.session_state["crew_list"]):
                        curr_row = 2 + row_idx 
                        ws.cell(row=curr_row, column=1).value = data["Name"]
                        ws.cell(row=curr_row, column=2).value = data["Surname"]
                        ws.cell(row=curr_row, column=3).value = data["Date of Birth"]
                        ws.cell(row=curr_row, column=4).value = data["Nationality"]
                        ws.cell(row=curr_row, column=5).value = data["Type Of Crew (REF)"]
                        ws.cell(row=curr_row, column=6).value = data["Document Type(REF)"]
                        ws.cell(row=curr_row, column=7).value = data["Document Number"]
                    out = io.BytesIO()
                    wb.save(out)
                    st.success(trans[curr_lang]["ok_match"].format(len(st.session_state["crew_list"])))
                    st.download_button(label=trans[curr_lang]["dl_lbl"], data=out.getvalue(), file_name="Iraqi_Airways_Crew_Report.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, key="download_btn_c")
                except Exception as e: st.error(f"Error: {e}")
    else: st.info(trans[curr_lang]["empty"])
